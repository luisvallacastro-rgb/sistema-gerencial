#!/usr/bin/env python3
"""Extract the legacy Excel inventory into KMI's neutral JSON seed format.

This script never edits the workbook or a KMI database. The Excel Timestamp
column is the canonical movement date because some visible Fecha cells use an
ambiguous locale conversion.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


def decimal_value(value, field: str) -> Decimal:
    try:
        number = Decimal(str(value if value not in (None, "") else 0))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise ValueError(f"{field} no es numerico: {value!r}") from error
    if not number.is_finite():
        raise ValueError(f"{field} no es finito: {value!r}")
    return number


def iso_timestamp(value) -> str:
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).isoformat()
    raise ValueError(f"Timestamp invalido: {value!r}")


def clean(value) -> str:
    return " ".join(str(value or "").split())


def scaled_int(value: Decimal, scale: int) -> int:
    return int((value * scale).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def extract(source: Path) -> dict:
    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=True)

    catalog_sheet = workbook["CATALOGO_MP"]
    catalog_headers = [cell.value for cell in catalog_sheet[1]]
    catalog_index = {name: index for index, name in enumerate(catalog_headers)}
    items = []
    item_codes = set()
    for row in catalog_sheet.iter_rows(min_row=2, values_only=True):
        code = clean(row[catalog_index["CodigoMP"]]).upper()
        if not code:
            continue
        if code in item_codes:
            raise ValueError(f"Codigo de catalogo duplicado: {code}")
        item_codes.add(code)
        category = clean(row[catalog_index["Categoria"]])
        color = clean(row[catalog_index["Color"]])
        minimum = decimal_value(row[catalog_index["StockMinimo"]], f"StockMinimo {code}")
        notes = []
        if category:
            notes.append(f"Categoria: {category}")
        if color:
            notes.append(f"Color: {color}")
        if minimum:
            notes.append(f"Stock minimo: {minimum}")
        active_text = clean(row[catalog_index["Activo"]]).upper()
        items.append({
            "code": code,
            "description": clean(row[catalog_index["Descripcion"]]),
            "unit": clean(row[catalog_index["UnidadMedida"]]) or "Unidad",
            "notes": ". ".join(notes),
            "active": active_text != "NO",
        })

    movement_sheet = workbook["DB_MP_MOVIMIENTOS"]
    movement_headers = [cell.value for cell in movement_sheet[1]]
    movement_index = {name: index for index, name in enumerate(movement_headers)}
    movements = []
    skipped_zero = []
    seen_ids = set()
    for row in movement_sheet.iter_rows(min_row=2, values_only=True):
        source_id = int(row[movement_index["ID"]])
        if source_id in seen_ids:
            raise ValueError(f"ID de movimiento duplicado: {source_id}")
        seen_ids.add(source_id)
        code = clean(row[movement_index["CodigoMP"]]).upper()
        if code not in item_codes:
            raise ValueError(f"Movimiento {source_id} usa codigo inexistente: {code}")
        movement_type = clean(row[movement_index["TipoMovimiento"]]).upper()
        if movement_type not in {"ENTRADA", "SALIDA"}:
            raise ValueError(f"Movimiento {source_id} tiene tipo invalido: {movement_type}")
        quantity = decimal_value(row[movement_index["Cantidad"]], f"Cantidad movimiento {source_id}")
        if quantity < 0:
            raise ValueError(f"Movimiento {source_id} tiene cantidad negativa")
        if quantity == 0:
            skipped_zero.append(source_id)
            continue
        unit_cost = decimal_value(row[movement_index["CostoUnitario"]], f"Costo movimiento {source_id}")
        if unit_cost < 0:
            raise ValueError(f"Movimiento {source_id} tiene costo negativo")
        reference_parts = [
            clean(row[movement_index["DocumentoSoporte"]]),
            clean(row[movement_index["NoDocumento"]]),
            clean(row[movement_index["OP"]]),
        ]
        note_parts = [
            clean(row[movement_index["Proveedor"]]),
            clean(row[movement_index["Observaciones"]]),
        ]
        movements.append({
            "sourceId": source_id,
            "code": code,
            "date": iso_timestamp(row[movement_index["Timestamp"]]),
            "type": movement_type,
            "quantityMillis": scaled_int(quantity, 1000),
            "unitCostMicros": scaled_int(unit_cost, 1_000_000) if movement_type == "ENTRADA" else 0,
            "reference": " / ".join(part for part in reference_parts if part),
            "notes": " / ".join(part for part in note_parts if part),
            "createdBy": clean(row[movement_index["Usuario"]]) or "Importacion Excel",
        })

    movements.sort(key=lambda movement: (movement["date"], movement["sourceId"]))
    return {
        "schemaVersion": 1,
        "source": {
            "name": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "extractedAt": datetime.now().replace(microsecond=0).isoformat(),
            "datePolicy": "DB_MP_MOVIMIENTOS.Timestamp",
        },
        "items": items,
        "movements": movements,
        "audit": {
            "catalogRows": len(items),
            "sourceMovementRows": len(seen_ids),
            "importableMovementRows": len(movements),
            "skippedZeroQuantityRows": skipped_zero,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = extract(args.source.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload["audit"], ensure_ascii=False))


if __name__ == "__main__":
    main()
