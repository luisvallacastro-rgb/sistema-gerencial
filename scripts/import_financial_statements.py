#!/usr/bin/env python3
"""Convert the finance-owned workbook into the isolated statements snapshot."""

import argparse
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "outputs" / "sistema-gerencial" / "financial-statements-seed.json"


def clean_text(value):
    return " ".join(str(value or "").split())


def account_code(value, name):
    if value is None or not clean_text(value):
        calculated_rows = {
            "utilidad bruta": "UTILIDAD_BRUTA",
            "utilidad neta": "UTILIDAD_NETA",
        }
        return calculated_rows.get(clean_text(name).casefold(), "")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean_text(value)


def account_level(code):
    if not code.isdigit():
        return 0
    return {1: 0, 2: 1, 4: 2, 6: 3}.get(len(code), 4)


def numeric(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, str):
        value = value.replace("$", "").replace(",", "").strip()
    return round(float(value or 0), 2)


def first_amount(*values):
    numbers = [numeric(value) for value in values]
    return next((value for value in numbers if abs(value) > 0.000001), 0.0)


def parsed_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = clean_text(value)
    for pattern in ("%m/%d/%Y", "%m/%d/%y", "%b-%y", "%B-%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"Fecha financiera no reconocida: {raw!r}")


def balance_periods(worksheet):
    raw_dates = [parsed_date(row[0].value) for row in worksheet.iter_rows(min_row=3) if row[0].value not in (None, "")]
    dates_by_year = defaultdict(set)
    for value in raw_dates:
        dates_by_year[value.year].add((value.month, value.day))
    # In the supplied 2024/2025 balance history, Excel stores Jan-Dec as
    # 01/01, 01/02, ... 01/12. Detect that complete pattern by year instead
    # of treating every record as January.
    month_encoded_as_day = {
        year for year, markers in dates_by_year.items()
        if len(markers) >= 12 and all(month == 1 and 1 <= day <= 12 for month, day in markers)
    }

    def normalize(value):
        parsed = parsed_date(value)
        month = parsed.day if parsed.year in month_encoded_as_day and parsed.month == 1 else parsed.month
        return f"{parsed.year:04d}-{month:02d}"

    return normalize


def import_balance(worksheet):
    normalize_period = balance_periods(worksheet)
    rows = []
    for values in worksheet.iter_rows(min_row=3, values_only=True):
        raw_period, raw_code, raw_name, partial, account, total, section = values[:7]
        code = account_code(raw_code, raw_name)
        if not code or raw_period in (None, ""):
            continue
        rows.append({
            "period": normalize_period(raw_period),
            "code": code,
            "name": clean_text(raw_name),
            "level": account_level(code),
            "isAuxiliary": code.isdigit() and len(code) >= 8,
            "amount": first_amount(section, total, account, partial),
        })
    return rows


def import_income(worksheet):
    rows = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        raw_period, raw_code, raw_name, partial_period, total_period, partial_accumulated, total_accumulated = values[:7]
        code = account_code(raw_code, raw_name)
        if not code or raw_period in (None, ""):
            continue
        parsed = parsed_date(raw_period)
        rows.append({
            "period": f"{parsed.year:04d}-{parsed.month:02d}",
            "code": code,
            "name": clean_text(raw_name),
            "level": account_level(code),
            "isAuxiliary": code.isdigit() and len(code) >= 8,
            "periodAmount": first_amount(total_period, partial_period),
            "accumulatedAmount": first_amount(total_accumulated, partial_accumulated),
        })
    return rows


def validate_periods(balance, income):
    balance_periods_found = sorted({row["period"] for row in balance})
    income_periods_found = sorted({row["period"] for row in income})
    if balance_periods_found != income_periods_found:
        missing_balance = sorted(set(income_periods_found) - set(balance_periods_found))
        missing_income = sorted(set(balance_periods_found) - set(income_periods_found))
        raise ValueError(f"Cobertura mensual distinta. Balance faltante={missing_balance}; resultados faltante={missing_income}")
    if not balance_periods_found or not re.fullmatch(r"\d{4}-\d{2}", balance_periods_found[-1]):
        raise ValueError("No se detectó una cobertura mensual válida")
    return balance_periods_found


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    balance = import_balance(workbook["Balances"])
    income = import_income(workbook["Resultados"])
    periods = validate_periods(balance, income)
    payload = {
        "source": {
            "file": args.workbook.name,
            "generatedAt": datetime.now().date().isoformat(),
            "currency": "USD",
            "coverageFrom": periods[0],
            "coverageTo": periods[-1],
        },
        "periods": periods,
        "balance": balance,
        "income": income,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Períodos: {periods[0]} a {periods[-1]} ({len(periods)})")
    print(f"Balance: {len(balance)} filas; Resultados: {len(income)} filas")
    print(args.output)


if __name__ == "__main__":
    main()
